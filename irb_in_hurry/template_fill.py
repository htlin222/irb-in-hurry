"""Fill an official blank DOCX in place — the institution-agnostic generator.

Instead of rebuilding a form with python-docx (the KFSYSCC generators), copy
the institution's own blank and write values next to its labels. Layout,
margins, fonts and wording stay exactly the official ones, so the layout gate
passes by construction.

    from irb_in_hurry.template_fill import blank_generator

    FORM_REGISTRY = {
        "F01": ("新案申請書", "institutions.myhosp.forms", "generate_f01"),
    }
    generate_f01 = blank_generator("F01", fields={
        "計畫名稱": "{study.title_zh}",           # label → format string
        "計畫主持人": "{pi.name}",
        "預計收案人數": lambda c: f"{c['subjects']['planned_n']} 人",
    }, checks={
        "簡易審查": lambda c: c["study"]["review_type"] == "expedited",
    })

fields: label text → value. The value goes into the next table cell to the
        right of the label (or below it if the label is the last cell), or
        replaces the blank (＿＿ / ____ / ：) after the label in a paragraph.
checks: option text → bool. A ☐/□ right before the option becomes ■ when true.

Theme fonts in the blank are pinned to the profile's form font, so the copy
renders identically in Word for Windows and Mac.
"""
import os
import re
import shutil

from docx import Document

from irb_in_hurry.docx_utils import check, pin_form_font, strip_theme_fonts
from irb_in_hurry.institution import current

UNCHECKED = "□☐"
_BLANK_RE = re.compile(r"[＿_]{2,}|\s{3,}")
_WS_RE = re.compile(r"[\s　:：]+")


def _norm(text):
    return _WS_RE.sub("", text or "")


def _lookup(config, dotted):
    node = config
    for key in dotted.split("."):
        if isinstance(node, list):
            node = node[int(key)]
        elif isinstance(node, dict):
            node = node.get(key, "")
        else:
            return ""
    return "" if node is None else node


def resolve(value, config):
    """'{pi.name}（{pi.dept}）' / callable / literal → str."""
    if callable(value):
        return str(value(config))
    return re.sub(r"\{([\w.]+)\}", lambda m: str(_lookup(config, m.group(1))), str(value))


def _set_cell_text(cell, text):
    """Replace a cell's text, keeping the first run's formatting."""
    paras = cell.paragraphs
    runs = [r for p in paras for r in p.runs]
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        paras[0].add_run(text)


def _iter_tables(doc):
    stack = list(doc.tables)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                stack.extend(part.tables)
    while stack:
        tbl = stack.pop(0)
        yield tbl
        for row in tbl.rows:
            for cell in row.cells:
                stack.extend(cell.tables)


def _iter_paragraphs(doc):
    yield from doc.paragraphs
    for tbl in _iter_tables(doc):
        for row in tbl.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _fill_table_label(doc, label, value):
    want = _norm(label)
    for tbl in _iter_tables(doc):
        rows = tbl.rows
        for ri, row in enumerate(rows):
            cells = row.cells
            for ci, cell in enumerate(cells):
                if _norm(cell.text) != want:
                    continue
                # merged cells repeat; take the first distinct cell to the right
                target = next((c for c in cells[ci + 1:] if c._tc is not cell._tc), None)
                if target is None and ri + 1 < len(rows):
                    target = rows[ri + 1].cells[ci]
                if target is not None:
                    _set_cell_text(target, value)
                    return True
    return False


def _fill_paragraph_label(doc, label, value):
    for p in _iter_paragraphs(doc):
        for run in p.runs:
            idx = run.text.find(label)
            if idx < 0:
                continue
            head, tail = run.text[:idx + len(label)], run.text[idx + len(label):]
            m = _BLANK_RE.search(tail)
            if m:
                run.text = head + tail[:m.start()] + value + tail[m.end():]
            elif tail.startswith(("：", ":")):
                run.text = head + tail[0] + value + tail[1:]
            else:
                run.text = head + "：" + value + tail
            return True
    return False


def _set_check(doc, option, on):
    pattern = re.compile(rf"[{UNCHECKED}■☑](\s*){re.escape(option)}")
    hit = False
    for p in _iter_paragraphs(doc):
        runs = p.runs
        text = "".join(r.text for r in runs)
        for m in pattern.finditer(text):
            # Word splits runs freely: flip the box glyph in whichever run holds it
            pos = m.start()
            for run in runs:
                if pos < len(run.text):
                    run.text = run.text[:pos] + check(on) + run.text[pos + 1:]
                    break
                pos -= len(run.text)
            hit = True
    return hit


def fill_blank(blank_path, out_path, fields=None, checks=None, config=None):
    """Copy blank → out_path and fill it. Returns list of labels not found."""
    shutil.copyfile(blank_path, out_path)
    doc = Document(out_path)
    missing = []
    for label, value in (fields or {}).items():
        text = resolve(value, config or {})
        if not (_fill_table_label(doc, label, text) or _fill_paragraph_label(doc, label, text)):
            missing.append(label)
    for option, cond in (checks or {}).items():
        on = bool(cond(config) if callable(cond) else cond)
        if not _set_check(doc, option, on):
            missing.append(f"{check(True)}{option}")
    pin_form_font(doc)       # theme fonts render differently on Windows vs Mac
    strip_theme_fonts(doc)
    doc.save(out_path)
    return missing


def blank_generator(form_id, fields=None, checks=None, filename="{form_id}_{study.irb_no}.docx"):
    """Build a `generate(config, output_dir) -> path` for one official blank."""
    def generate(config, output_dir):
        blank = current().blank_path(form_id)
        if not os.path.exists(blank):
            raise FileNotFoundError(f"No blank for {form_id}: {blank} "
                                    f"(run `irbh templates`, or map it in templates.files)")
        name = resolve(filename.replace("{form_id}", form_id), config) or f"{form_id}.docx"
        out = os.path.join(output_dir, name)
        missing = fill_blank(blank, out, fields, checks, config)
        if missing:
            print(f"    ⚠ {form_id}: labels not found in blank: {', '.join(missing)}")
        return out
    generate.__name__ = f"generate_{form_id.lower()}"
    return generate
