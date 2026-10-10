"""Institution-agnostic pipeline: profile, fill-the-blank engine, onboarding."""
import os
import tomllib

from docx import Document
from docx.shared import Twips

from irb_in_hurry import institution, onboard
from irb_in_hurry.template_fill import fill_blank, resolve

CONFIG = {
    "study": {"irb_no": "2026-001", "title_zh": "示範研究", "review_type": "expedited"},
    "pi": {"name": "王小明", "phone": "1234"},
    "subjects": {"planned_n": 120, "consent_waiver": True},
}


def make_blank(path, margins=1000):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Twips(11906), Twips(16838)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Twips(margins)
    doc.add_paragraph("示範醫院 研究倫理委員會")
    doc.add_paragraph("新案申請書")
    tbl = doc.add_table(rows=3, cols=2)
    for i, label in enumerate(["IRB編號", "計畫名稱", "計畫主持人"]):
        tbl.cell(i, 0).text = label
    p = doc.add_paragraph()
    p.add_run("審查類別：□")            # box and option in different runs
    p.add_run("簡易審查　□一般審查")
    doc.add_paragraph("預計收案人數：＿＿＿＿ 人")
    doc.save(path)


def test_kfsyscc_is_default_profile():
    prof = institution.load("kfsyscc")
    assert prof.heading == "和信治癌中心醫院 人體試驗委員會"
    assert prof.irb_no_label == "KFSYSCC-IRB編號"
    assert prof.form_id("IRB_SF90 藥品試驗") == "SF090"
    assert prof.form_id("中文計畫摘要") == "PROPOSAL"
    assert institution.resolve_id({"institution": "other"}) in ("other", os.environ.get("IRB_INSTITUTION"))


def test_resolve_format_strings_and_callables():
    assert resolve("{pi.name}（{pi.phone}）", CONFIG) == "王小明（1234）"
    assert resolve(lambda c: c["subjects"]["planned_n"] * 2, CONFIG) == "240"
    assert resolve("{missing.key}", CONFIG) == ""


def test_fill_blank_labels_lines_and_checkboxes(tmp_path):
    blank, out = tmp_path / "blank.docx", tmp_path / "out.docx"
    make_blank(blank)
    missing = fill_blank(str(blank), str(out), fields={
        "IRB編號": "{study.irb_no}",
        "計畫主持人": "{pi.name}",
        "預計收案人數": "{subjects.planned_n}",
        "不存在的欄位": "x",
    }, checks={
        "簡易審查": lambda c: c["study"]["review_type"] == "expedited",
        "一般審查": False,
    }, config=CONFIG)

    doc = Document(str(out))
    cells = {r.cells[0].text: r.cells[1].text for r in doc.tables[0].rows}
    text = "\n".join(p.text for p in doc.paragraphs)
    assert cells["IRB編號"] == "2026-001" and cells["計畫主持人"] == "王小明"
    assert "■簡易審查" in text and "□一般審查" in text
    assert "預計收案人數：120 人" in text
    assert missing == ["不存在的欄位"]
    # page setup of the official blank is untouched
    assert doc.sections[0].left_margin == Twips(1000)


def test_onboard_drafts_profile_forms_and_inventory(tmp_path):
    tpl = tmp_path / "templates" / "demo"
    tpl.mkdir(parents=True)
    make_blank(tpl / "IRB-F1 新案申請書.docx")
    make_blank(tpl / "IRB-F2 檢核表.docx", margins=800)

    assert onboard.main("demo", root=str(tmp_path)) == 0
    out = tmp_path / "institutions" / "demo"
    prof = tomllib.loads((out / "profile.toml").read_text(encoding="utf-8"))
    assert prof["forms_module"] == "institutions.demo.forms"
    assert prof["name"] == "示範醫院" and prof["committee"] == "研究倫理委員會"
    assert prof["templates"]["files"] == {"F001": "IRB-F1 新案申請書.docx", "F002": "IRB-F2 檢核表.docx"}
    assert prof["page"]["margins"]["left"] in (1000, 800)
    assert len(prof["page"]["per_form_margins"]) == 1

    forms_py = (out / "forms.py").read_text(encoding="utf-8")
    compile(forms_py, "forms.py", "exec")
    assert "'計畫主持人': \"{pi.name}\"" in forms_py
    assert "'簡易審查': lambda c:" in forms_py
    assert "FORM_REGISTRY" in forms_py and "PHASE_FORMS" in forms_py
    assert "預計收案人數" in (out / "form_inventory.md").read_text(encoding="utf-8")
