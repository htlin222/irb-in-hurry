"""End-to-end tests: config → selected forms → DOCX files."""
import glob
import os

import pytest
from docx import Document

from scripts import generate_all
from scripts.docx_utils import load_config
from scripts.form_selector import FORM_REGISTRY, PHASE_FORMS, get_generator, select_forms

FIXTURES = sorted(glob.glob("tests/fixtures/*.yml"))


@pytest.fixture
def retro_config():
    return load_config("tests/fixtures/sample_retrospective.yml")


@pytest.fixture
def output_dir(tmp_path):
    path = tmp_path / "output"
    path.mkdir()
    return str(path)


def docx_text(path):
    """All paragraph and table-cell text of a DOCX."""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def generate_phase(config, phase, output_dir):
    config["phase"] = phase
    paths = []
    for fid, _ in select_forms(config):
        path = generate_all.generate_form(fid, config, output_dir)
        assert os.path.getsize(path) > 0, f"{fid} file is empty: {path}"
        paths.append(path)
    return paths


@pytest.mark.parametrize("form_id", sorted(FORM_REGISTRY))
def test_every_registered_generator_runs(form_id, retro_config, output_dir):
    """Each registry entry points at an importable function that writes a DOCX."""
    assert get_generator(form_id) is not None
    path = generate_all.generate_form(form_id, retro_config, output_dir)
    assert os.path.exists(path)


@pytest.mark.parametrize("phase", sorted(PHASE_FORMS))
@pytest.mark.parametrize("fixture", FIXTURES, ids=os.path.basename)
def test_generate_all_every_fixture_and_phase(fixture, phase, tmp_path):
    """The orchestrator succeeds for every shipped fixture in every phase."""
    out = tmp_path / "output"
    code = generate_all.main(fixture, str(out), phase=phase, checklist_path=str(tmp_path / "checklist.md"))
    assert code == 0
    assert len(list(out.glob("*.docx"))) == len(select_forms(load_config(fixture) | {"phase": phase}))


def test_new_case_generates_all_forms(retro_config, output_dir):
    paths = generate_phase(retro_config, "new", output_dir)
    assert len(paths) == 6  # SF001, SF002, SF094, PROPOSAL, SF003, SF005


def test_closure_generates_all_forms(retro_config, output_dir):
    paths = generate_phase(retro_config, "closure", output_dir)
    assert len(paths) == 4  # SF036, SF037, SF038, SF023


def test_generate_all_reports_bad_config(tmp_path, capsys):
    bad = tmp_path / "config.yml"
    bad.write_text("phase: new\nstudy: {}\npi: {}\n", encoding="utf-8")
    assert generate_all.main(str(bad), str(tmp_path / "out")) == 2
    assert "study.title_zh" in capsys.readouterr().out
    assert generate_all.main(str(tmp_path / "missing.yml"), str(tmp_path / "out")) == 2


def test_docx_contains_irb_number(retro_config, output_dir):
    path = generate_all.generate_form("SF001", retro_config, output_dir)
    assert "20250801A" in docx_text(path), "IRB number not found in SF001"


def test_docx_contains_pi_name(retro_config, output_dir):
    path = generate_all.generate_form("SF002", retro_config, output_dir)
    assert "林協霆" in docx_text(path), "PI name not found in SF002"


def test_checklist_generation(retro_config, tmp_path):
    """Verify checklist.md is generated correctly."""
    from scripts.checklist import generate_checklist

    results = [
        ("SF001", "新案審查送審資料表", "/fake/path.docx", "generated"),
        ("SF002", "研究計畫申請書", "/fake/path2.docx", "generated"),
        ("SF094", "顯著財務利益申報表", None, "error"),
    ]
    checklist_path = str(tmp_path / "checklist.md")
    generate_checklist(retro_config, results, "新案審查", checklist_path)

    with open(checklist_path) as f:
        content = f.read()

    assert "20250801A" in content
    assert "■ SF001" in content  # generated
    assert "■ SF002" in content  # generated
    assert "□ SF094" in content  # error
    assert "irb@kfsyscc.org" in content


def test_config_validation():
    """Verify config loads without error."""
    config = load_config("tests/fixtures/sample_retrospective.yml")
    assert config["study"]["irb_no"] == "20250801A"
    assert config["pi"]["name"] == "林協霆"
    assert config["subjects"]["consent_waiver"] is True
    assert len(config["co_pi"]) == 1
    assert config["co_pi"][0]["name"] == "邱倫維"


def test_proposal_summary_uses_config_text(output_dir):
    """proposal.* text fills 中文計畫摘要; absent keys keep the placeholder."""
    from scripts.generators.proposal import generate_proposal_summary

    config = load_config("tests/fixtures/example_tdxd_her2low.yml")
    text = "\n".join(p.text for p in Document(
        generate_proposal_summary(config, output_dir)).paragraphs)
    assert "DESTINY-Breast04" in text
    assert "1. 主要目的" in text
    assert "（請列出）" not in text
    assert "標準。疾病" in text  # folded-line spaces removed

    config.pop("proposal")
    text = "\n".join(p.text for p in Document(
        generate_proposal_summary(config, output_dir)).paragraphs)
    assert "納入條件：（請列出）" in text
