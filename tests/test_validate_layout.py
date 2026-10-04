"""Layout safety gate: generated forms must be A4 + 標楷體 on every platform."""
import importlib

import pytest
from docx import Document

from scripts.docx_utils import (
    add_p,
    apply_official_page_setup,
    form_filename,
    form_id_from_path,
    init_doc,
    load_config,
    official_margins,
)
from scripts.form_selector import get_generator, select_forms
from scripts.validate_layout import validate_file

PHASES = ["new", "amendment", "continuing", "closure"]


@pytest.fixture
def retro_config():
    return load_config("tests/fixtures/sample_retrospective.yml")


def _generate(config, phase, out):
    config["phase"] = phase
    paths = []
    for fid, _ in select_forms(config):
        mod_path, func_name = get_generator(fid)
        path = getattr(importlib.import_module(f"scripts.{mod_path}"), func_name)(config, out)
        apply_official_page_setup(path)
        paths.append(path)
    return paths


@pytest.mark.parametrize("phase", PHASES)
def test_generated_forms_pass_layout_gate(retro_config, tmp_path, phase):
    for path in _generate(retro_config, phase, str(tmp_path)):
        rep = validate_file(path, render=False)
        assert rep.errors == [], f"{rep.name}: {rep.errors}"
        assert not any("邊界" in w for w in rep.warnings), f"{rep.name}: {rep.warnings}"


def test_generated_form_is_a4_with_official_margins(retro_config, tmp_path):
    path = next(p for p in _generate(retro_config, "new", str(tmp_path)) if "SF002" in p)
    sect = Document(path).sections[0]
    assert (sect.page_width.twips, sect.page_height.twips) == (11906, 16838)
    margins = tuple(m.twips for m in (sect.left_margin, sect.right_margin,
                                      sect.top_margin, sect.bottom_margin))
    assert margins == official_margins("SF002")


def test_init_doc_declares_kai_font_for_win_and_mac(tmp_path):
    doc = init_doc()
    path = str(tmp_path / "SF001_x.docx")
    doc.save(path)
    blob = next(p for p in Document(path).part.package.iter_parts()
                if str(p.partname) == "/word/fontTable.xml").blob.decode()
    assert 'w:name="標楷體"' in blob and "DFKai-SB" in blob


def test_plain_python_docx_is_rejected(tmp_path):
    """The old failure mode: US Letter page + CJK text on theme fonts."""
    doc = Document()
    doc.add_paragraph("計畫名稱")
    path = str(tmp_path / "SF001_plain.docx")
    doc.save(path)
    rep = validate_file(path, render=False)
    assert any("紙張大小" in e for e in rep.errors)
    assert any("標楷體" in e for e in rep.errors)


def test_non_big5_characters_warned(tmp_path):
    doc = init_doc()
    add_p(doc, "酪胺酸激酶抑制劑")   # 酶 is outside Big5
    add_p(doc, "简体字")
    path = str(tmp_path / "SF001_glyph.docx")
    doc.save(path)
    apply_official_page_setup(path)
    rep = validate_file(path, render=False)
    assert rep.errors == []
    assert any("Big5" in w and "简" in w for w in rep.warnings)


@pytest.mark.parametrize("name,fid", [
    ("SF002_KF-001.docx", "SF002"),
    ("IRB_SF90_同意書.docx", "SF090"),
    ("中文計畫摘要_proposal.docx", "PROPOSAL"),
    ("中文計畫摘要_20250801A.docx", "PROPOSAL"),  # the name proposal.py actually writes
    ("checklist.docx", None),
])
def test_form_id_from_path(name, fid):
    assert form_id_from_path(name) == fid


@pytest.mark.parametrize("irb_no,expected", [
    ("20250801A", "SF001_20250801A.docx"),
    ("", "SF001_新案審查送審資料表.docx"),
    ("KF/2025 01", "SF001_KF-2025-01.docx"),
])
def test_form_filename(irb_no, expected):
    assert form_filename("SF001", {"study": {"irb_no": irb_no}}, "新案審查送審資料表") == expected
