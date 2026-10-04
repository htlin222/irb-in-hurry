"""Shared DOCX helper functions for IRB form generation.

Extracted from irb-close/generate_forms.py, refactored to accept config dict.
"""
import os
import re
import yaml
from docx import Document
from docx.shared import Pt, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def load_config(path="config.yml"):
    """Load and return config dict from YAML file."""
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def check(condition: bool) -> str:
    """Return ■ if True, □ if False (IRB checkbox convention)."""
    return "■" if condition else "□"


# Official KFSYSCC blank forms: A4 portrait, L/R 2 cm, T/B 2.5 cm, header 1.5 cm.
# (python-docx's default template is US Letter with 3.17 cm sides, which shifts
# every line break and table width relative to the official form.)
# Values are in twips, copied from the official templates' sectPr.
PAGE_SETUP = {
    "page_width": Twips(11906), "page_height": Twips(16838),
    "left_margin": Twips(1134), "right_margin": Twips(1134),
    "top_margin": Twips(1418), "bottom_margin": Twips(1418),
    "header_distance": Twips(851), "footer_distance": Twips(992),
}

# Forms whose official blank uses other margins: (left, right, top, bottom) twips.
# Extracted from the official templates (make templates); all others use PAGE_SETUP.
OFFICIAL_MARGINS = {
    "PROPOSAL": (1134, 1134, 1134, 1134),
    "SF002": (1134, 1134, 1418, 1134),
    "SF023": (1134, 1134, 1418, 1134),
    "SF031": (1134, 1134, 1418, 1134),
    "SF032": (1134, 1134, 1418, 567),
    "SF047": (1134, 1134, 1418, 1134),
    "SF062": (851, 851, 776, 800),
    "SF063": (680, 737, 851, 1134),
    "SF066": (1134, 1134, 1418, 851),
    "SF067": (1134, 1134, 907, 284),
    "SF068": (1134, 1134, 907, 1048),
    "SF075": (680, 737, 851, 1134),
    "SF079": (1134, 1134, 1418, 1134),
    "SF080": (1134, 1134, 1418, 1134),
    "SF082": (1134, 1134, 1418, 851),
    "SF083": (1134, 1134, 1134, 1418),
    "SF084": (1134, 1134, 1418, 851),
    "SF085": (1134, 1134, 1134, 1418),
    "SF090": (680, 737, 851, 1134),
    "SF091": (851, 851, 851, 1134),
    "SF092": (851, 851, 851, 1134),
    "SF094": (1134, 1134, 1134, 1134),
}
_FORM_ID_RE = re.compile(r"SF\s*0*(\d{1,3})")


def form_id_from_path(path):
    """'SF002_KF-001.docx' → 'SF002'; '中文計畫摘要_proposal.docx' → 'PROPOSAL'."""
    name = os.path.basename(path)
    m = _FORM_ID_RE.search(name)
    if m:
        return f"SF{int(m.group(1)):03d}"
    return "PROPOSAL" if "proposal" in name.lower() else None


def official_margins(form_id):
    """(left, right, top, bottom) in twips of the official blank form."""
    if form_id in OFFICIAL_MARGINS:
        return OFFICIAL_MARGINS[form_id]
    return tuple(int(PAGE_SETUP[k].twips) for k in
                 ("left_margin", "right_margin", "top_margin", "bottom_margin"))


def apply_official_page_setup(path):
    """Re-save a generated DOCX with A4 + the official margins of its form."""
    doc = Document(path)
    left, right, top, bottom = official_margins(form_id_from_path(path))
    for section in doc.sections:
        section.page_width, section.page_height = PAGE_SETUP["page_width"], PAGE_SETUP["page_height"]
        section.left_margin, section.right_margin = Twips(left), Twips(right)
        section.top_margin, section.bottom_margin = Twips(top), Twips(bottom)
    doc.save(path)


FORM_FONT = '標楷體'
# fontTable entry Word itself writes for 標楷體. altName lets Word resolve the
# font by its English name (DFKai-SB on Windows); macOS ships it as BiauKai
# with the same localized name, so both platforms find the real Kai font.
_FONT_TABLE_ENTRY = (
    '<w:font w:name="標楷體"><w:altName w:val="DFKai-SB"/>'
    '<w:panose1 w:val="03000509000000000000"/><w:charset w:val="88"/>'
    '<w:family w:val="script"/><w:pitch w:val="fixed"/>'
    '<w:sig w:usb0="00000003" w:usb1="080E0000" w:usb2="00000016" w:usb3="00000000"'
    ' w:csb0="00100001" w:csb1="00000000"/></w:font>'
)


def _apply_cross_platform_defaults(doc):
    """Pin page setup, fonts and language so Word (Win/Mac) and LibreOffice agree."""
    for section in doc.sections:
        for attr, value in PAGE_SETUP.items():
            setattr(section, attr, value)

    # docDefaults: replace theme fonts (which resolve to Calibri/MS 明朝 on an
    # en-US theme) with explicit 標楷體, and tag text as Traditional Chinese.
    rpr_default = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    fonts = rpr_default.find(qn('w:rFonts'))
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        fonts.attrib.pop(qn(a), None)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(a), FORM_FONT)
    lang = rpr_default.find(qn('w:lang'))
    lang.set(qn('w:eastAsia'), 'zh-TW')

    for part in doc.part.package.iter_parts():
        if str(part.partname) == '/word/fontTable.xml' and FORM_FONT.encode() not in part.blob:
            part._blob = part.blob.replace(
                b'</w:fonts>', _FONT_TABLE_ENTRY.encode('utf-8') + b'</w:fonts>')


def init_doc(sz=12):
    """Create a new A4 Document with 標楷體 default font."""
    doc = Document()
    _apply_cross_platform_defaults(doc)
    s = doc.styles['Normal']
    s.font.name = FORM_FONT
    s.font.size = Pt(sz)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), FORM_FONT)
    return doc


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, **kwargs):
    """Set cell borders. kwargs: top, bottom, start, end with {val, sz, color}."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, attrs in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{attrs.get("val", "single")}" '
            f'w:sz="{attrs.get("sz", 4)}" w:space="0" '
            f'w:color="{attrs.get("color", "000000")}"/>'
        )
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_run_font(run, font_name="標楷體", size=12, bold=False):
    """Set run font with eastAsia fallback."""
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run.font.element.rPr
    if rPr is None:
        run.font.element.get_or_add_rPr()
        rPr = run.font.element.rPr
    rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_p(doc, text, bold=False, size=12, alignment=None, sa=Pt(6), sb=Pt(0)):
    """Add a formatted paragraph to the document."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, "標楷體", size, bold)
    if alignment:
        p.alignment = alignment
    p.paragraph_format.space_after = sa
    p.paragraph_format.space_before = sb
    return p


def add_ct(cell, text, bold=False, size=10, alignment=None):
    """Add formatted text to a table cell."""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.clear()
    run = p.add_run(text)
    set_run_font(run, "標楷體", size, bold)
    if alignment:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    return p


def apply_tb(tbl):
    """Apply borders to all cells in a table."""
    for row in tbl.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"sz": 4}, bottom={"sz": 4},
                start={"sz": 4}, end={"sz": 4},
            )


def add_header(doc, config, inc_proj=True):
    """Add standard IRB header block (IRB no, title, PI, dept)."""
    rows = 5 if inc_proj else 4
    tbl = doc.add_table(rows=rows, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    i = 0
    add_ct(tbl.rows[i].cells[0], "KFSYSCC-IRB編號", True, 11)
    add_ct(tbl.rows[i].cells[1], config["study"]["irb_no"], size=11)

    if inc_proj:
        i += 1
        add_ct(tbl.rows[i].cells[0], "計畫編號", True, 11)
        add_ct(tbl.rows[i].cells[1], config["study"].get("project_no", "不適用"), size=11)

    i += 1
    add_ct(tbl.rows[i].cells[0], "計畫名稱", True, 11)
    cell = tbl.rows[i].cells[1]
    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(f"（中文）{config['study']['title_zh']}")
    set_run_font(run, size=10)
    p.paragraph_format.space_after = Pt(2)
    p2 = cell.add_paragraph()
    run2 = p2.add_run(f"（英文）{config['study']['title_en']}")
    set_run_font(run2, size=10)
    p2.paragraph_format.space_after = Pt(2)

    i += 1
    # Build PI/co-PI string
    pi_str = config["pi"]["name"]
    co_pis = config.get("co_pi", [])
    if co_pis:
        pi_str += "（主持人）"
        for cp in co_pis:
            pi_str += f"\n{cp['name']}（共同主持人）"
    add_ct(tbl.rows[i].cells[0], "計畫主持人", True, 11)
    add_ct(tbl.rows[i].cells[1], pi_str, size=11)

    i += 1
    add_ct(tbl.rows[i].cells[0], "單位／職稱", True, 11)
    add_ct(tbl.rows[i].cells[1], config["pi"]["dept"], size=11)

    apply_tb(tbl)
    doc.add_paragraph()


def add_footer(doc, ver, fno, date):
    """Add version/form number footer."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"版次第{ver}版　{fno}　{date}")
    set_run_font(run, size=9)
