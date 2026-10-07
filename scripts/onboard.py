#!/usr/bin/env python3
"""Onboard a new institution from its blank DOCX forms.

Put the institution's official blank forms (.docx / .doc) in
templates/<id>/ — any filenames — then:

    make onboard INST=<id>     # or: uv run python scripts/onboard.py <id>

It reads every blank and writes a *draft* pipeline to institutions/<id>/:

    profile.toml        page size, default + per-form margins, form font,
                        institution/committee name and IRB-number label guesses
    forms.py            FORM_REGISTRY + PHASE_FORMS using the generic
                        fill-the-blank generator, labels pre-mapped to config
                        fields where recognised, the rest left as TODOs
    form_inventory.md   per form: title, page setup, fonts, table labels,
                        checkbox options, fill-in blanks — the worksheet you
                        (or Claude) use to finish forms.py

Nothing is overwritten unless --force. Review every `# TODO` before use.
"""
import collections
import json
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

FORM_EXTS = (".docx", ".doc")
CODE_RE = re.compile(r"([A-Za-z]{1,5})[-_ .]?0*(\d{1,4})")
BOX_RE = re.compile(r"[□☐■☑]\s*([^□☐■☑\s：:，,、。（）()]{1,20})")
BLANK_RE = re.compile(r"([^\s＿_：:]{1,12})\s*[：:]?\s*(?:[＿_]{2,})")
CJK_RE = re.compile(r"[㐀-鿿]")

# Label keywords → config field (first match wins). Extend for your forms.
LABEL_HINTS = [
    (("IRB", "編號"), "{study.irb_no}"),
    (("IRB", "No"), "{study.irb_no}"),
    (("計畫編號",), "{study.project_no}"),
    (("英文",), "{study.title_en}"),
    (("計畫名稱",), "{study.title_zh}"),
    (("研究名稱",), "{study.title_zh}"),
    (("Protocol", "Title"), "{study.title_en}"),
    (("共同主持人",), "{co_pi.0.name}"),
    (("主持人",), "{pi.name}"),
    (("Investigator",), "{pi.name_en}"),
    (("單位",), "{pi.dept}"),
    (("職稱",), "{pi.dept}"),
    (("電話",), "{pi.phone}"),
    (("分機",), "{pi.phone}"),
    (("E-mail",), "{pi.email}"),
    (("Email",), "{pi.email}"),
    (("電子郵件",), "{pi.email}"),
    (("人數",), "{subjects.planned_n}"),
    (("執行期間",), "{dates.study_start} 至 {dates.study_end}"),
    (("起始",), "{dates.study_start}"),
    (("結束",), "{dates.study_end}"),
]
CHECK_HINTS = [
    ("簡易審查", 'lambda c: c["study"].get("review_type") == "expedited"'),
    ("免予審查", 'lambda c: c["study"].get("review_type") == "exempt"'),
    ("一般審查", 'lambda c: c["study"].get("review_type") == "full_board"'),
    ("免除知情同意", 'lambda c: c["subjects"].get("consent_waiver")'),
    ("免取得知情同意", 'lambda c: c["subjects"].get("consent_waiver")'),
    ("回溯", 'lambda c: c["study"].get("type") == "retrospective"'),
    ("前瞻", 'lambda c: c["study"].get("type") == "prospective"'),
    ("多中心", 'lambda c: c["study"].get("multicenter")'),
    ("單一中心", 'lambda c: not c["study"].get("multicenter")'),
    ("基因", 'lambda c: c["study"].get("genetic")'),
]


def to_docx(path):
    if path.endswith(".docx"):
        return path
    from scripts.convert import find_soffice
    soffice = find_soffice()
    target = os.path.splitext(path)[0] + ".docx"
    if soffice and not os.path.exists(target):
        subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir",
                        os.path.dirname(path), path], capture_output=True, timeout=180)
    return target if os.path.exists(target) else None


def assign_ids(files):
    """Stable form ids: the code in the filename if every file has one, else F01…"""
    stems = [os.path.splitext(f)[0] for f in files]
    codes = [CODE_RE.search(s) for s in stems]
    prefixes = {m.group(1).upper() for m in codes if m}
    if all(codes) and len(prefixes) == 1:
        prefix = prefixes.pop()
        width = max(3, max(len(m.group(2)) for m in codes))
        ids = [f"{prefix}{int(m.group(2)):0{width}d}" for m in codes]
        if len(set(ids)) == len(ids):
            return ids, rf"{prefix}[-_ .]?0*(\d{{1,4}})", f"{prefix}{{:0{width}d}}"
    return [f"F{i:02d}" for i in range(1, len(files) + 1)], r"F(\d{2})", "F{:02d}"


def inspect(path):
    doc = Document(path)
    sect = doc.sections[0]
    def twips(length):
        return round(int(length or 0) / 635)   # EMU → twips

    info = {
        "width": twips(sect.page_width), "height": twips(sect.page_height),
        "margins": tuple(twips(getattr(sect, k)) for k in
                         ("left_margin", "right_margin", "top_margin", "bottom_margin")),
    }

    fonts, latin = collections.Counter(), collections.Counter()
    for r in doc.element.body.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        if rf is not None and text.strip():
            if CJK_RE.search(text) and rf.get(qn("w:eastAsia")):
                fonts[rf.get(qn("w:eastAsia"))] += len(text)
            if rf.get(qn("w:ascii")):
                latin[rf.get(qn("w:ascii"))] += len(text)
    dd = doc.styles.element.find(qn("w:docDefaults"))
    default_rf = dd.find(".//" + qn("w:rFonts")) if dd is not None else None
    if not fonts and default_rf is not None and default_rf.get(qn("w:eastAsia")):
        fonts[default_rf.get(qn("w:eastAsia"))] += 1
    lang = dd.find(".//" + qn("w:lang")) if dd is not None else None
    info["font"] = fonts.most_common(1)[0][0] if fonts else None
    info["latin_font"] = latin.most_common(1)[0][0] if latin else None
    info["lang"] = lang.get(qn("w:eastAsia")) if lang is not None else None

    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    info["title_lines"] = paras[:3]

    labels, seen = [], set()
    for tbl in doc.tables:
        for row in tbl.rows:
            cells, prev = row.cells, None
            for i, c in enumerate(cells):
                if prev is not None and c._tc is prev._tc:
                    continue
                prev = c
                t = " ".join(c.text.split())
                nxt = next((x for x in cells[i + 1:] if x._tc is not c._tc), None)
                if t and len(t) <= 20 and not BOX_RE.match(t) and t not in seen \
                        and (nxt is None or not nxt.text.strip()):
                    seen.add(t)
                    labels.append(t)
    info["labels"] = labels

    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for tbl in doc.tables for row in tbl.rows for c in row.cells)
    info["has_cjk"] = bool(CJK_RE.search(all_text))
    info["checks"] = list(dict.fromkeys(m.group(1) for m in BOX_RE.finditer(all_text)))
    info["blanks"] = list(dict.fromkeys(m.group(1) for m in BLANK_RE.finditer(all_text)))
    info["irb_label"] = next((lbl for lbl in labels if "IRB" in lbl.upper() and
                              any(k in lbl for k in ("編號", "No", "NO", "number"))), None)
    return info


def guess_field(label):
    for keys, field in LABEL_HINTS:
        if all(k.lower() in label.lower() for k in keys):
            return field
    return None


def guess_check(option):
    return next((expr for key, expr in CHECK_HINTS if key in option), None)


def form_title(info, fallback):
    lines = info["title_lines"]
    # 1st line is usually the institution heading, 2nd the form name
    return (lines[1] if len(lines) > 1 else lines[0]) if lines else fallback


def write_profile(inst_id, forms, id_rule, path, force=False):
    def common(key):
        return collections.Counter(f[key] for f in forms.values() if f[key]).most_common(1)

    size = common("width"), common("height")
    width = size[0][0][0] if size[0] else 11906
    height = size[1][0][0] if size[1] else 16838
    margins = common("margins")[0][0] if common("margins") else (1134, 1134, 1418, 1418)
    font = common("font")[0][0] if common("font") else "標楷體"
    lang = common("lang")[0][0] if common("lang") else None
    if not lang or (lang[:2] not in ("zh", "ja", "ko") and any(f["has_cjk"] for f in forms.values())):
        lang = "zh-TW"   # blank had no East Asian language tag; CJK text found
    heading = common_heading(forms)
    name, committee = (heading.split(" ", 1) + [""])[:2] if heading else ("TODO 醫院名稱", "人體試驗委員會")
    irb_label = common("irb_label")[0][0] if common("irb_label") else "IRB編號"
    per_form = {fid: list(f["margins"]) for fid, f in forms.items()
                if f["margins"] != margins and any(f["margins"])}
    aliases = {"標楷體": ["DFKai-SB", "BiauKai"], "新細明體": ["PMingLiU"], "細明體": ["MingLiU"],
               "微軟正黑體": ["Microsoft JhengHei"], "宋体": ["SimSun"], "Times New Roman": []}
    encoding = {"zh-TW": "cp950", "zh-CN": "gbk", "ja-JP": "cp932", "ko-KR": "cp949"}.get(lang)

    lines = [
        f"# Draft profile generated by scripts/onboard.py from templates/{inst_id}/.",
        f'# Review every TODO, then set `institution = "{inst_id}"` in config.toml.',
        f"id = {_q(inst_id)}",
        f"name = {_q(name)}            # TODO confirm",
        'name_en = "TODO"',
        f"committee = {_q(committee or '人體試驗委員會')}   # TODO confirm",
        f"irb_no_label = {_q(irb_label)}",
        f"ehr_name = {_q(name + '電子病歷系統')}",
        "",
        f"forms_module = {_q(f'institutions.{inst_id}.forms')}",
        "",
        "[submission]",
        'email = "TODO@example.org"',
        'paper = "TODO (e.g. 1 original + 1 copy to IRB office)"',
        "",
        "[templates]",
        '# index_url = "https://irb.example.org/forms?page={n}"   # optional scraper',
        "# index_range = [1, 3]",
        '# link_host = "direct"          # direct | google_drive',
        f"form_id_pattern = {_q(id_rule[0])}",
        f"form_id_format = {_q(id_rule[1])}",
        "named_forms = {}",
        "max_pages = {}               # forms allowed to grow past their blank, e.g. { F02 = 3 }",
        "",
        f"[templates.files]            # form id → blank in templates/{inst_id}/",
        *[f"{fid} = {_q(f['blank'])}" for fid, f in forms.items()],
        "",
        "[page]",
        f"width = {width}",
        f"height = {height}",
        "margins = { left = %d, right = %d, top = %d, bottom = %d }" % margins,
        "header_distance = 851",
        "footer_distance = 992",
        "",
        "[page.per_form_margins]      # [left, right, top, bottom] where a blank differs",
    ]
    lines += [f"{fid} = {m}" for fid, m in sorted(per_form.items())]
    lines += [
        "",
        "[font]",
        f"name = {_q(font)}",
        f"aliases = {_q(aliases.get(font, []))}   # TODO Windows/macOS names of the face",
        f"lang = {_q(lang)}                 # TODO verify (w:lang eastAsia)",
        f"encoding = {_q(encoding)}" if encoding else "# encoding = \"cp950\"     # glyphs the font is guaranteed to have",
        "",
    ]
    _write(path, "\n".join(lines), force)


def common_heading(forms):
    firsts = collections.Counter(f["title_lines"][0] for f in forms.values() if f["title_lines"])
    return firsts.most_common(1)[0][0] if firsts else None


def write_forms_py(inst_id, forms, path, force=False):
    out = [
        f'"""{inst_id} form pack — DRAFT generated by scripts/onboard.py.',
        "",
        "Each form fills the institution's own blank (templates/<id>/<FORM>.docx)",
        "via scripts.template_fill.blank_generator. Finish the TODOs using",
        "form_inventory.md, then define real routing in PHASE_FORMS.",
        '"""',
        "from scripts.template_fill import blank_generator",
        "",
    ]
    reg = []
    for fid, f in forms.items():
        title = form_title(f, fid)
        func = f"generate_{fid.lower()}"
        reg.append(f'    "{fid}": ({title!r}, "institutions.{inst_id}.forms", "{func}"),')
        out.append(f"# {fid} — {title}  (source: {f['source']})")
        out.append(f'{func} = blank_generator("{fid}", fields={{')
        for label in f["labels"]:
            field = guess_field(label)
            out.append(f'    {label!r}: "{field}",' if field else f'    # {label!r}: "",  # TODO')
        for blank in f["blanks"]:
            if blank not in f["labels"]:
                field = guess_field(blank)
                out.append(f'    {blank!r}: "{field}",' if field else f'    # {blank!r}: "",  # TODO (fill-in line)')
        out.append("}, checks={")
        for opt in f["checks"]:
            expr = guess_check(opt)
            out.append(f'    {opt!r}: {expr},' if expr else f'    # {opt!r}: lambda c: False,  # TODO')
        out += ["})", ""]

    out += ["", "FORM_REGISTRY = {", *reg, "}", "",
            "# TODO: split forms into the phases your IRB uses (new, amendment,",
            "# continuing, closure, sae, …) and add conditional forms, e.g.",
            '#   (lambda c: c["subjects"]["consent_waiver"], ["F05"]),',
            "PHASE_FORMS = {",
            '    "new": {',
            f'        "base": {list(forms)},',
            '        "conditions": [],',
            "    },",
            "}", ""]
    _write(path, "\n".join(out), force)


def write_inventory(inst_id, forms, path, force=False):
    out = [f"# Form inventory — {inst_id}", "",
           "Generated by `scripts/onboard.py`. One section per official blank. Use it to",
           "finish `forms.py`: map every label / fill-in line to a config field, give",
           "every checkbox a condition, and group the forms into phases.", "",
           "| Form | Title | Paper (cm) | Margins L/R/T/B (cm) | Font |",
           "|------|-------|-----------|----------------------|------|"]
    def cm(v):
        return f"{v / 567:.2f}"

    for fid, f in forms.items():
        out.append(f"| {fid} | {form_title(f, fid)} | {cm(f['width'])}×{cm(f['height'])} | "
                   f"{'/'.join(cm(m) for m in f['margins'])} | {f['font'] or '?'} |")
    for fid, f in forms.items():
        out += ["", f"## {fid} — {form_title(f, fid)}", "", f"- Source file: `{f['source']}`",
                f"- Heading lines: {' / '.join(f['title_lines']) or '—'}",
                f"- Fonts: CJK `{f['font']}`, Latin `{f['latin_font']}`, lang `{f['lang']}`", "",
                "**Table labels** (empty cell to the right → value goes there)", ""]
        out += [f"- `{lbl}` → {guess_field(lbl) or '**TODO**'}" for lbl in f["labels"]] or ["- —"]
        out += ["", "**Fill-in lines** (`label：＿＿＿`)", ""]
        out += [f"- `{b}` → {guess_field(b) or '**TODO**'}" for b in f["blanks"]] or ["- —"]
        out += ["", "**Checkbox options**", ""]
        out += [f"- □ {c} → {guess_check(c) or '**TODO**'}" for c in f["checks"]] or ["- —"]
    _write(path, "\n".join(out) + "\n", force)


def _q(value):
    """TOML literal for a str or list of str (JSON strings are valid TOML basic strings)."""
    return json.dumps(value, ensure_ascii=False)


def _write(path, text, force):
    if os.path.exists(path) and not force:
        print(f"  · keep {os.path.relpath(path, ROOT)} (exists; --force to overwrite)")
        return
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
    print(f"  ■ wrote {os.path.relpath(path, ROOT)}")


def list_blanks(tpl_dir):
    """One file per form: a .doc and its converted .docx count once (prefer .docx)."""
    by_stem = {}
    for f in sorted(os.listdir(tpl_dir)):
        stem, ext = os.path.splitext(f)
        if ext.lower() in FORM_EXTS and not f.startswith(("~$", ".")):
            if stem not in by_stem or ext.lower() == ".docx":
                by_stem[stem] = f
    return sorted(by_stem.values())


def main(inst_id, root=ROOT, force=False):
    tpl_dir = os.path.join(root, "templates", inst_id)
    if not os.path.isdir(tpl_dir):
        print(f"Put the blank forms in templates/{inst_id}/ first")
        return 1
    files = list_blanks(tpl_dir)
    if not files:
        print(f"No .docx/.doc files in templates/{inst_id}/")
        return 1

    ids, *id_rule = assign_ids(files)
    forms = {}
    print(f"Inspecting {len(files)} blank forms in templates/{inst_id}/ ...")
    for fid, f in zip(ids, files, strict=True):
        docx = to_docx(os.path.join(tpl_dir, f))
        if not docx:
            print(f"  ✗ {f}: .doc needs LibreOffice to convert — skipped")
            continue
        info = inspect(docx)
        info["source"], info["blank"] = f, os.path.basename(docx)
        forms[fid] = info
        print(f"  ■ {fid} ← {f}  ({len(info['labels'])} labels, {len(info['checks'])} checkboxes)")

    out_dir = os.path.join(root, "institutions", inst_id)
    os.makedirs(out_dir, exist_ok=True)
    init = os.path.join(out_dir, "__init__.py")
    if not os.path.exists(init):
        open(init, "w").close()
    write_profile(inst_id, forms, id_rule, os.path.join(out_dir, "profile.toml"), force)
    write_forms_py(inst_id, forms, os.path.join(out_dir, "forms.py"), force)
    write_inventory(inst_id, forms, os.path.join(out_dir, "form_inventory.md"), force)
    print(f"\nNext: finish the TODOs in institutions/{inst_id}/ (docs/ONBOARDING.md), set "
          f'`institution = "{inst_id}"` in config.toml, then `make all`.')
    return 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(args[0], force="--force" in sys.argv))
